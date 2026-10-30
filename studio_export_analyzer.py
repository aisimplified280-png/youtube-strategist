"""Parse and analyse YouTube Studio CSV/TSV exports without persisting uploaded files.

The parser intentionally recognizes only known metric families. Unknown columns are reported
as ignored rather than guessed into a metric, and all analysis stays in the current session.
"""
from __future__ import annotations

import csv
import io
import math
import re
from collections import defaultdict
from datetime import date, datetime
from statistics import median


HEADER_ALIASES = {
    "video_title": {
        "video title", "content title", "title", "video name", "content", "video",
        "video headline", "asset", "item",
    },
    "video_id": {
        "video id", "content id", "video identifier", "videoid", "youtube video id",
    },
    "date": {
        "date", "day", "report date", "date published", "video publish time",
        "video published time", "publish date",
    },
    "views": {"views", "engaged views", "organic views", "view count", "total views"},
    "watch_time": {
        "watch time hours", "watch time (hours)", "watch time minutes",
        "watch time (minutes)", "estimated minutes watched",
        "estimated minutes watched (minutes)", "watch time", "watch hours",
    },
    "average_view_duration": {
        "average view duration", "avg view duration", "average view duration seconds",
        "average view duration (seconds)", "avg duration", "average duration",
    },
    "average_view_percentage": {
        "average percentage viewed", "average view percentage",
        "average percentage viewed (%)", "average view duration percentage",
        "avg percentage viewed", "avg view %", "avg view percentage",
        "percent viewed", "percentage viewed", "average view percentage (%)",
    },
    "impressions": {
        "impressions", "thumbnail impressions", "video thumbnail impressions", "total impressions",
    },
    "ctr": {
        "impressions click through rate", "impressions click-through rate",
        "click through rate", "click-through rate", "video thumbnail impressions ctr",
        "thumbnail click through rate", "ctr", "impressions ctr", "thumbnail ctr",
        "impressions click-through rate (%)", "impressions click through rate (%)", "ctr (%)",
    },
    "subscribers_gained": {
        "subscribers gained", "subscribers", "subscribers added", "new subscribers",
    },
    "subscribers_lost": {"subscribers lost"},
    "likes": {"likes"},
    "comments": {"comments", "comment count", "comments added", "total comments"},
    "shares": {"shares", "share count"},
    "traffic_source": {
        "traffic source", "traffic source type", "youtube traffic source",
        "insight traffic source type",
    },
    "search_term": {"search term", "top search terms", "youtube search term", "search query"},
    "viewer_percentage": {"viewer percentage", "viewer percentage of total"},
    "geography": {"geography", "country", "country code", "region"},
    "age_group": {"age group", "viewer age", "age"},
    "gender": {"gender", "viewer gender"},
    "device_type": {"device type", "device"},
    "elapsed_time": {
        "elapsed video time", "elapsed time", "video time", "elapsed video time (seconds)",
    },
    "retention_percent": {
        "audience retention", "audience retention percentage", "retention rate",
        "absolute audience retention",
    },
    "relative_retention": {"relative audience retention", "relative retention"},
}

NUMERIC_FIELDS = {
    "views", "watch_time", "average_view_duration", "average_view_percentage",
    "impressions", "ctr", "subscribers_gained", "subscribers_lost",
    "likes", "comments", "shares", "retention_percent", "relative_retention",
}


MAX_UPLOAD_BYTES = 10 * 1024 * 1024


class StudioExportError(ValueError):
    """Raised when an upload is unsupported or has no recognizable report header."""


def safe_csv_export_value(value):
    """Neutralize spreadsheet formula prefixes in user-visible text before CSV download."""
    if not isinstance(value, str):
        return value
    if value.lstrip(" \t\r\n").startswith(("=", "+", "-", "@")):
        return "'" + value
    return value

def _norm(value: str) -> str:
    value = (value or "").replace("\ufeff", "").strip().casefold()
    value = value.replace("&", " and ")
    value = re.sub(r"[_:/()%-]+", " ", value)
    return re.sub(r"\s+", " ", value).strip()


_NORMALIZED_ALIASES = {
    field: {_norm(alias) for alias in aliases}
    for field, aliases in HEADER_ALIASES.items()
}


def _map_header(value: str):
    normalized = _norm(value)
    if not normalized:
        return None
    for field, aliases in _NORMALIZED_ALIASES.items():
        if normalized in aliases:
            return field
    # Common export variants where the label includes a unit or explanatory suffix.
    if normalized.startswith("average view duration") or normalized.startswith("avg view duration"):
        return "average_view_duration"
    if (
        normalized.startswith("average percentage viewed")
        or normalized.startswith("average view percentage")
        or normalized.startswith("avg percentage viewed")
        or normalized.startswith("avg view percentage")
        or normalized.startswith("percent viewed")
        or normalized.startswith("percentage viewed")
    ):
        return "average_view_percentage"
    if (
        normalized.startswith("impressions click through rate")
        or normalized.startswith("video thumbnail impressions ctr")
        or normalized.startswith("impressions ctr")
        or normalized == "ctr"
        or "ctr" in normalized.split()
    ):
        return "ctr"
    if normalized.startswith("video thumbnail impressions") or normalized == "impressions":
        return "impressions"
    if normalized.startswith("estimated minutes watched"):
        return "watch_time"
    if normalized.startswith("watch time") and "hour" in normalized:
        return "watch_time"
    if normalized.startswith("watch time") and "minute" in normalized:
        return "watch_time"
    if "relative audience retention" in normalized:
        return "relative_retention"
    if "audience retention" in normalized:
        return "retention_percent"
    if normalized.startswith("elapsed video time"):
        return "elapsed_time"
    if normalized in {"subscribers", "subscribers added", "subscribers gained"}:
        return "subscribers_gained"
    if normalized in {"comments added", "comments"}:
        return "comments"
    return None


def _map_row_columns(row: list) -> dict[str, int]:
    """Map a header row to known metric fields, disambiguating multi-identifier exports."""
    normalized_labels = [_norm(str(label or "")) for label in row]

    has_content = any(norm == "content" for norm in normalized_labels)
    has_video_title = any(norm in {"video title", "title", "video name"} for norm in normalized_labels)
    has_video = any(norm == "video" for norm in normalized_labels)

    mapped = {}
    for column_index, label in enumerate(row):
        norm = normalized_labels[column_index]
        if has_content and has_video_title and norm == "content":
            field = "video_id"
        elif has_video and has_video_title and norm == "video":
            field = "video_id"
        else:
            field = _map_header(label)
        if field and field not in mapped:
            mapped[field] = column_index
    return mapped


def _parse_number(value):
    if value is None:
        return None
    text = str(value).strip()
    if not text or text.casefold() in {"—", "-", "n/a", "na", "not available", "(not set)"}:
        return None
    text = text.replace(",", "").replace("\u2212", "-")
    if text.startswith("(") and text.endswith(")"):
        text = "-" + text[1:-1]
    text = text.replace("%", "").strip()
    match = re.search(r"[-+]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][-+]?\d+)?", text)
    if not match:
        return None
    try:
        result = float(match.group(0))
    except ValueError:
        return None
    return result if math.isfinite(result) else None


def _parse_duration_seconds(value):
    text = str(value or "").strip()
    if not text:
        return None
    # Studio may display average duration as M:SS or H:MM:SS.
    if re.fullmatch(r"\d{1,3}:\d{2}(?::\d{2})?", text):
        chunks = [int(part) for part in text.split(":")]
        if len(chunks) == 2:
            return chunks[0] * 60 + chunks[1]
        return chunks[0] * 3600 + chunks[1] * 60 + chunks[2]
    return _parse_number(text)


def _parse_elapsed(value):
    text = str(value or "").strip()
    if not text:
        return None
    if text.endswith("%"):
        number = _parse_number(text)
        return {"seconds": None, "percent": number}
    return {"seconds": _parse_duration_seconds(text), "percent": None}


def _is_percent_label(header):
    return "%" in (header or "") or "percent" in (header or "").casefold() or "rate" in (header or "").casefold()


def _to_percent(value, raw, header):
    if value is None:
        return None
    if "%" in str(raw):
        return value
    # If the column header explicitly specifies percentage/rate with % or percent/percentage,
    # Studio exports the values already in percentage points (e.g. 0.8 means 0.8%, not 80%).
    header_str = str(header or "").casefold()
    if "%" in header_str or "percentage" in header_str or "percent" in header_str:
        return value
    # Studio may export raw decimal ratios (e.g., 0.054 for 5.4%) in sheets without %;
    # only normalize ratios strictly in [0, 1] when the header has no percent notation.
    if 0 <= value <= 1:
        return value * 100
    return value


def _detect_delimiter(text):
    sample = "\n".join(text.splitlines()[:80])
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",\t;")
        return dialect.delimiter
    except csv.Error:
        candidates = [",", "\t", ";"]
        lines = text.splitlines()[:80]
        scores = {}
        for delimiter in candidates:
            counts = [line.count(delimiter) for line in lines if line.strip()]
            scores[delimiter] = (sum(count > 0 for count in counts), sum(counts))
        return max(candidates, key=lambda item: scores[item])


def _decode_upload(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-16", "cp1252", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise StudioExportError("Could not decode this file. Export it as CSV or TSV and upload again.")


def _excel_cell_text(value, number_format=""):
    """Convert an Excel cell into text without losing dates, numeric values, or percents."""
    if value is None:
        return ""
    if isinstance(value, (int, float)) and "%" in (number_format or ""):
        return f"{value * 100:g}%"
    if isinstance(value, datetime):
        if value.time().hour == 0 and value.time().minute == 0 and value.time().second == 0:
            return value.date().isoformat()
        return value.isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def _read_xlsx_matrix(data: bytes) -> tuple[list[list[str]], str]:
    """Read the worksheet that most clearly contains a Studio report table."""
    try:
        from openpyxl import load_workbook
    except ImportError as error:
        raise StudioExportError(
            "Excel workbook support is unavailable. Upload CSV/TSV or contact the app owner."
        ) from error
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as error:
        raise StudioExportError(
            "This does not look like a valid .xlsx workbook. Open it in Excel and save it as .xlsx, then try again."
        ) from error

    best_matrix = None
    best_sheet = None
    best_rank = None
    for worksheet in workbook.worksheets:
        matrix = [
            [_excel_cell_text(cell.value, cell.number_format) for cell in row]
            for row in worksheet.iter_rows()
        ]
        best_header = None
        for row_index, row in enumerate(matrix[:60]):
            mapped = _map_row_columns(row)
            if len(mapped) < 2:
                continue
            important = sum(field in mapped for field in (
                "video_title", "views", "impressions", "ctr",
                "elapsed_time", "retention_percent", "relative_retention",
            ))
            rank = (len(mapped) * 10 + important, -row_index)
            if best_header is None or rank > best_header:
                best_header = rank
        if best_header is None:
            continue
        rank = (best_header[0], best_header[1], len(matrix))
        if best_rank is None or rank > best_rank:
            best_rank = rank
            best_matrix = matrix
            best_sheet = worksheet.title
    workbook.close()

    if best_matrix is None:
        raise StudioExportError(
            "I couldn't find a YouTube Studio report table in this workbook. Choose a worksheet with video titles/IDs and measurable report columns."
        )
    return best_matrix, best_sheet


def parse_studio_export(filename: str, data: bytes) -> dict:
    """Parse one Studio CSV/TSV/TXT or Excel workbook export in memory."""
    suffix = (filename or "").lower().rsplit(".", 1)[-1] if "." in (filename or "") else ""
    if suffix not in {"csv", "tsv", "txt", "xlsx"}:
        raise StudioExportError(
            "Upload a YouTube Studio CSV, TSV, TXT, or .xlsx workbook."
        )
    if len(data) > MAX_UPLOAD_BYTES:
        raise StudioExportError(
            "This file exceeds the 10 MB per-file limit. Export a narrower date range or split the report."
        )

    source_sheet = None
    if suffix == "xlsx":
        matrix, source_sheet = _read_xlsx_matrix(data)
        delimiter = "excel"
    else:
        text = _decode_upload(data)
        if not text.strip():
            raise StudioExportError("This file is empty.")
        delimiter = "\t" if suffix == "tsv" else _detect_delimiter(text)
        try:
            matrix = list(csv.reader(io.StringIO(text), delimiter=delimiter))
        except csv.Error as error:
            raise StudioExportError("The file could not be parsed as a CSV/TSV report.") from error
    if not matrix:
        raise StudioExportError("This file contains no rows.")

    header_candidates = []
    for index, row in enumerate(matrix[:60]):
        mapped = _map_row_columns(row)
        if len(mapped) >= 2:
            # Prefer a complete report header and break ties by reaching one later in preamble.
            important = sum(field in mapped for field in (
                "video_title", "views", "impressions", "ctr",
                "elapsed_time", "retention_percent", "relative_retention",
            ))
            header_candidates.append((len(mapped) * 10 + important, index, mapped, row))
    if not header_candidates:
        raise StudioExportError(
            "I couldn't identify a YouTube Studio report header. Upload an export from "
            "Analytics → Advanced mode / See more → Export current view."
        )
    _, header_index, columns, header_row = max(header_candidates, key=lambda item: (item[0], item[1]))
    header_labels = {field: header_row[index].strip() for field, index in columns.items()}
    records = []
    for row in matrix[header_index + 1:]:
        if not any(str(cell).strip() for cell in row):
            continue
        if len(row) <= max(columns.values(), default=-1):
            row = row + [""] * (max(columns.values()) + 1 - len(row))
        # Ignore repeated header lines and completely unmeasured table footers.
        if any(_map_header(cell) == field and str(cell).strip()
               for field, index in columns.items() for cell in [row[index]]):
            continue
        record = {}
        for field, index in columns.items():
            raw = str(row[index]).strip() if index < len(row) else ""
            label = header_labels.get(field, "")
            if field == "video_title":
                if raw and re.fullmatch(r"[A-Za-z0-9_-]{11}", raw):
                    if not record.get("video_id"):
                        record["video_id"] = raw
                    record["video_title"] = None
                else:
                    record["video_title"] = raw or None
            elif field == "video_id":
                if raw:
                    match = re.search(r"(?:v=|youtu\.be/|/shorts/|/embed/)([A-Za-z0-9_-]{11})", raw)
                    if re.fullmatch(r"[A-Za-z0-9_-]{11}", raw):
                        record["video_id"] = raw
                    elif match:
                        record["video_id"] = match.group(1)
                    else:
                        record["video_id"] = raw
                else:
                    record["video_id"] = None
            elif field in {"date", "traffic_source", "search_term",
                         "geography", "age_group", "gender", "device_type"}:
                record[field] = raw or None
            elif field == "elapsed_time":
                elapsed = _parse_elapsed(raw)
                record["elapsed_seconds"] = elapsed["seconds"]
                record["elapsed_percent"] = elapsed["percent"]
                record["elapsed_label"] = raw or None
            elif field == "watch_time":
                number = _parse_number(raw)
                normalized_label = _norm(label)
                if number is not None and "hour" in normalized_label:
                    number *= 60
                record["watch_time_minutes"] = number
            elif field == "average_view_duration":
                record["average_view_duration_seconds"] = _parse_duration_seconds(raw)
            elif field == "average_view_percentage":
                number = _parse_number(raw)
                record["average_view_percentage"] = _to_percent(number, raw, label)
            elif field == "ctr":
                number = _parse_number(raw)
                record["ctr_percent"] = _to_percent(number, raw, label)
            elif field == "retention_percent":
                number = _parse_number(raw)
                record["retention_percent"] = _to_percent(number, raw, label)
            elif field == "relative_retention":
                record["relative_retention_index"] = _parse_number(raw)
            else:
                number = _parse_number(raw)
                if field in {"views", "impressions", "subscribers_gained", "subscribers_lost",
                             "likes", "comments", "shares"} and number is not None:
                    number = int(number) if number.is_integer() else number
                record[field] = number
        if any(value is not None for key, value in record.items() if key not in {"video_title", "date", "traffic_source"}):
            records.append(record)

    if not records:
        raise StudioExportError("A report header was found, but no rows with usable measurements followed it.")

    fields = set(columns)
    if "elapsed_time" in fields and ({"retention_percent", "relative_retention"} & fields):
        report_type = "retention"
    elif ({"video_title", "video_id"} & fields) and ({"views", "impressions", "ctr", "watch_time"} & fields):
        report_type = "video_performance"
    elif "traffic_source" in fields:
        report_type = "traffic_sources"
    elif {"search_term", "geography", "age_group", "gender", "device_type"} & fields:
        report_type = "breakdown"
    elif "date" in fields and ({"views", "watch_time"} & fields):
        report_type = "daily_trend"
    else:
        report_type = "analytics_table"

    return {
        "filename": filename,
        "report_type": report_type,
        "delimiter": "excel" if suffix == "xlsx" else ("tab" if delimiter == "\t" else ("semicolon" if delimiter == ";" else "comma")),
        "source_sheet": source_sheet,
        "header_row": header_index + 1,
        "header_labels": header_labels,
        "fields": sorted(fields),
        "ignored_headers": [
            str(label).strip() for label in header_row
            if str(label).strip() and _map_header(label) is None
        ],
        "records": records,
    }


def _weighted_average(rows, value_key, weight_key):
    values = [(row.get(value_key), row.get(weight_key)) for row in rows]
    values = [(value, weight if isinstance(weight, (int, float)) and weight > 0 else 1)
              for value, weight in values if isinstance(value, (int, float))]
    if not values:
        return None
    total_weight = sum(weight for _, weight in values)
    return sum(value * weight for value, weight in values) / total_weight if total_weight else None


def _video_summaries(records):
    groups = defaultdict(list)
    for row in records:
        title = row.get("video_title")
        video_id = row.get("video_id")
        if not title and not video_id:
            continue
        if title and title.strip().casefold() in {"total", "overall", "all videos", "grand total"}:
            continue
        key = ("id", video_id) if video_id else ("title", _norm(title))
        groups[key].append(row)

    result = []
    for (key_type, key_value), rows in groups.items():
        first = rows[0]
        valid_title = None
        for r in rows:
            t = r.get("video_title")
            if t and not re.fullmatch(r"[A-Za-z0-9_-]{11}", str(t).strip()) and "://" not in str(t):
                valid_title = str(t).strip()
                break
        valid_id = None
        for r in rows:
            vid = r.get("video_id")
            if vid and re.fullmatch(r"[A-Za-z0-9_-]{11}", str(vid).strip()):
                valid_id = str(vid).strip()
                break
        if not valid_id:
            for r in rows:
                t = r.get("video_title")
                if t and re.fullmatch(r"[A-Za-z0-9_-]{11}", str(t).strip()):
                    valid_id = str(t).strip()
                    break

        item = {
            "video_id": valid_id or first.get("video_id"),
            "video_title": valid_title,
            "row_count": len(rows),
        }
        for field in ("views", "watch_time_minutes", "impressions", "subscribers_gained",
                      "subscribers_lost", "likes", "comments", "shares"):
            available = [row.get(field) for row in rows if isinstance(row.get(field), (int, float))]
            item[field] = sum(available) if available else None
        item["ctr_percent"] = _weighted_average(rows, "ctr_percent", "impressions")
        item["average_view_duration_seconds"] = _weighted_average(
            rows, "average_view_duration_seconds", "views"
        )
        item["average_view_percentage"] = _weighted_average(
            rows, "average_view_percentage", "views"
        )
        date_values = [row.get("date") for row in rows if row.get("date")]
        item["date_start"] = min(date_values) if date_values else None
        item["date_end"] = max(date_values) if date_values else None
        result.append(item)
    return result


def _elapsed_sort_key(row):
    if isinstance(row.get("elapsed_seconds"), (int, float)):
        return row["elapsed_seconds"]
    if isinstance(row.get("elapsed_percent"), (int, float)):
        return row["elapsed_percent"] * 10000
    return float("inf")


def _retention_recommendations(records):
    recommendations = []
    absolute = sorted(
        [row for row in records if isinstance(row.get("retention_percent"), (int, float))],
        key=_elapsed_sort_key,
    )
    if len(absolute) >= 3:
        changes = []
        for previous, current in zip(absolute, absolute[1:]):
            drop = previous["retention_percent"] - current["retention_percent"]
            if drop > 0:
                changes.append((drop, previous, current))
        for drop, previous, current in sorted(changes, key=lambda item: item[0], reverse=True)[:3]:
            if drop < 3:
                continue
            start = previous.get("elapsed_label") or "previous point"
            end = current.get("elapsed_label") or "next point"
            recommendations.append({
                "priority": "Review",
                "title": f"Inspect the retention drop between {start} and {end}",
                "evidence": f"Reported audience retention decreased by {drop:.1f} percentage points between adjacent exported points.",
                "action": "Review the visual/narration transition and whether it matches the promise of the opening. The curve shows where to inspect, not why viewers left.",
            })
        if not recommendations:
            recommendations.append({
                "priority": "Monitor",
                "title": "No large adjacent retention drop surfaced in this export",
                "evidence": "The parsed absolute-retention points do not contain an adjacent drop of 3 percentage points or more.",
                "action": "Compare the curve against the opening promise and inspect the raw report for smaller dips or spikes.",
            })
    relative = [row for row in records if isinstance(row.get("relative_retention_index"), (int, float))]
    if relative:
        recommendations.append({
            "priority": "Context",
            "title": "Relative retention is a comparison index, not audience-retained percentage",
            "evidence": "This export contains relative audience-retention values.",
            "action": "Use the relative line to compare interest at equivalent moments, but do not interpret its values as the exact percent of viewers still watching.",
        })
    return recommendations


def _format_video_label(title: str | None) -> str:
    val = str(title or "").strip()
    if val and not re.fullmatch(r"[A-Za-z0-9_-]{11}", val) and "://" not in val:
        return val
    return "Video title unavailable"


def _performance_recommendations(videos):
    recommendations = []

    # 1. CTR Recommendations: qualify videos with impressions >= 50 and valid CTR
    ctr_rows = [
        row for row in videos
        if isinstance(row.get("ctr_percent"), (int, float))
        and isinstance(row.get("impressions"), (int, float))
        and row["impressions"] >= 50
    ]
    if len(ctr_rows) >= 2:
        midpoint = median(row["ctr_percent"] for row in ctr_rows)
        low_ctr = sorted(
            [row for row in ctr_rows if row["ctr_percent"] < midpoint],
            key=lambda row: (row["ctr_percent"], -row["impressions"]),
        )[:5]
        for row in low_ctr[:3]:
            video_label = _format_video_label(row.get("video_title"))
            recommendations.append({
                "priority": "High" if row["impressions"] >= 1000 else "Review",
                "title": f"Review title/thumbnail alignment: {video_label}",
                "evidence": (
                    f"Reported CTR {row['ctr_percent']:.2f}% is below this upload set's "
                    f"median ({midpoint:.2f}%) across {len(ctr_rows)} videos with at least 50 impressions."
                ),
                "action": "Check whether the title and thumbnail communicate one clear, accurate reason to click. Treat this as a review candidate, not proof the packaging caused the difference.",
                "video_id": row.get("video_id"),
                "video_title": row.get("video_title") if row.get("video_title") != "Video title unavailable" else None,
                "measured_metrics": {
                    "ctr_percent": row.get("ctr_percent"),
                    "impressions": row.get("impressions"),
                    "median_ctr_percent": midpoint,
                    "views": row.get("views"),
                },
            })
    elif len(ctr_rows) == 1:
        single = ctr_rows[0]
        if single["ctr_percent"] < 3.0:
            video_label = _format_video_label(single.get("video_title"))
            recommendations.append({
                "priority": "Review",
                "title": f"Review title/thumbnail alignment: {video_label}",
                "evidence": (
                    f"Reported CTR {single['ctr_percent']:.2f}% across {single['impressions']:,.0f} impressions "
                    "is below the typical 3.0% threshold."
                ),
                "action": "Check whether the title and thumbnail communicate one clear, accurate reason to click. Treat this as a review candidate, not proof the packaging caused the difference.",
                "video_id": single.get("video_id"),
                "video_title": single.get("video_title") if single.get("video_title") != "Video title unavailable" else None,
                "measured_metrics": {
                    "ctr_percent": single.get("ctr_percent"),
                    "impressions": single.get("impressions"),
                    "views": single.get("views"),
                },
            })

    # 2. Retention Recommendations: qualify videos with numeric average view percentage
    retention_rows = [
        row for row in videos
        if isinstance(row.get("average_view_percentage"), (int, float))
    ]
    if len(retention_rows) >= 2:
        midpoint = median(row["average_view_percentage"] for row in retention_rows)
        lowest = sorted(retention_rows, key=lambda row: row["average_view_percentage"])[:3]
        for row in lowest:
            if row["average_view_percentage"] >= midpoint:
                continue
            video_label = _format_video_label(row.get("video_title"))
            recommendations.append({
                "priority": "Review",
                "title": f"Review opening and pacing: {video_label}",
                "evidence": (
                    f"Average percentage viewed is {row['average_view_percentage']:.1f}%, "
                    f"below this export's median ({midpoint:.1f}%) across {len(retention_rows)} videos."
                ),
                "action": "Compare the opening promise with how quickly the video delivers value. Account for differences in video length and format before drawing conclusions.",
                "video_id": row.get("video_id"),
                "video_title": row.get("video_title") if row.get("video_title") != "Video title unavailable" else None,
                "measured_metrics": {
                    "views": row.get("views"),
                    "average_view_percentage": row.get("average_view_percentage"),
                    "median_average_view_percentage": midpoint,
                    "average_view_duration_seconds": row.get("average_view_duration_seconds"),
                },
            })

    # 3. View Retention Opportunity Benchmarks
    view_retention_rows = [
        row for row in videos
        if isinstance(row.get("views"), (int, float))
        and isinstance(row.get("average_view_percentage"), (int, float))
    ]
    if len(view_retention_rows) >= 2:
        views_midpoint = median(row["views"] for row in view_retention_rows)
        retention_midpoint = median(
            row["average_view_percentage"] for row in view_retention_rows
        )
        winners = [
            row for row in view_retention_rows
            if row["views"] >= views_midpoint
            and row["average_view_percentage"] >= retention_midpoint
        ]
        for row in sorted(
            winners,
            key=lambda item: (
                item["views"] / max(views_midpoint, 1)
                + item["average_view_percentage"] / max(retention_midpoint, 1),
            ),
            reverse=True,
        )[:2]:
            video_label = _format_video_label(row.get("video_title"))
            recommendations.append({
                "priority": "Opportunity",
                "title": f"Study this stronger-viewing example: {video_label}",
                "evidence": (
                    f"{row['views']:,.0f} reported views and "
                    f"{row['average_view_percentage']:.1f}% average viewed; "
                    f"the upload-set medians are {views_midpoint:,.0f} views and "
                    f"{retention_midpoint:.1f}% average viewed."
                ),
                "action": (
                    "Inspect the topic promise, opening structure, and traffic sources for "
                    "patterns worth testing again. These metrics highlight a candidate to study; "
                    "they do not prove what caused its performance."
                ),
                "video_id": row.get("video_id"),
                "video_title": row.get("video_title") if row.get("video_title") != "Video title unavailable" else None,
                "measured_metrics": {
                    "views": row.get("views"),
                    "average_view_percentage": row.get("average_view_percentage"),
                    "median_views": views_midpoint,
                    "median_average_view_percentage": retention_midpoint,
                },
            })

    # 4. Views & Watch-Time Tier: active when Views are present (e.g. standard Content Table data.csv)
    views_rows = [
        row for row in videos
        if isinstance(row.get("views"), (int, float)) and row["views"] > 0
    ]
    existing_rec_vids = {r.get("video_id") for r in recommendations if r.get("video_id")}
    existing_rec_titles = {
        _norm(r.get("video_title")) for r in recommendations if r.get("video_title")
    }

    if len(views_rows) >= 3 and not recommendations:
        views_midpoint = median(row["views"] for row in views_rows)
        # Strongest performers as opportunity benchmarks
        strong = sorted(
            [row for row in views_rows if row["views"] >= views_midpoint],
            key=lambda r: -r["views"],
        )[:2]
        for row in strong:
            video_label = _format_video_label(row.get("video_title"))
            recommendations.append({
                "priority": "Opportunity",
                "title": f"Study this stronger-reach upload: {video_label}",
                "evidence": (
                    f"{row['views']:,.0f} reported views; upload-set median is "
                    f"{views_midpoint:,.0f} views across {len(views_rows)} videos with reported traffic."
                ),
                "action": (
                    "Inspect the topic angle, thumbnail style, and title structure for "
                    "patterns that resonated with your audience. This identifies a stronger-reach "
                    "video to study, not proof of cause."
                ),
                "video_id": row.get("video_id"),
                "video_title": row.get("video_title") if row.get("video_title") != "Video title unavailable" else None,
                "measured_metrics": {
                    "views": row.get("views"),
                    "median_views": views_midpoint,
                    "watch_time_minutes": row.get("watch_time_minutes"),
                },
            })

        # Underperforming reach for metadata / packaging review
        low_views = sorted(
            [row for row in views_rows if row["views"] < views_midpoint * 0.6],
            key=lambda r: r["views"],
        )[:3]
        for row in low_views:
            vid = row.get("video_id")
            vtitle = row.get("video_title")
            if vid and vid in existing_rec_vids:
                continue
            if vtitle and _norm(vtitle) in existing_rec_titles:
                continue
            video_label = _format_video_label(vtitle)
            recommendations.append({
                "priority": "Review",
                "title": f"Review discoverability and packaging: {video_label}",
                "evidence": (
                    f"{row['views']:,.0f} reported views is below this upload set's "
                    f"median ({views_midpoint:,.0f} views) across {len(views_rows)} videos with traffic."
                ),
                "action": (
                    "Review the title and description for search clarity and alignment with what viewers search for. "
                    "Treat this as a review candidate, not proof that metadata caused the difference."
                ),
                "video_id": vid,
                "video_title": vtitle if vtitle != "Video title unavailable" else None,
                "measured_metrics": {
                    "views": row.get("views"),
                    "median_views": views_midpoint,
                    "watch_time_minutes": row.get("watch_time_minutes"),
                },
            })

    if not recommendations:
        evidence_parts = [f"{len(videos)} video row(s) detected in this report."]
        if ctr_rows:
            evidence_parts.append(f"{len(ctr_rows)} video(s) met the 50-impression threshold with CTR.")
        else:
            evidence_parts.append("0 videos had both CTR and >= 50 impressions.")
        if retention_rows:
            evidence_parts.append(f"{len(retention_rows)} video(s) had Average percentage viewed.")
        else:
            evidence_parts.append("0 videos had Average percentage viewed.")
        evidence_str = " ".join(evidence_parts)

        recommendations.append({
            "priority": "Context",
            "title": "Not enough qualifying video evidence for a ranked action",
            "evidence": evidence_str,
            "action": (
                "To enable packaging/CTR recommendations, export with 'Impressions' and 'Impressions click-through rate (%)' "
                "(Analytics → Advanced mode → Reach). To enable retention recommendations, export with 'Average percentage viewed (%)'."
            ),
        })
    return recommendations


def _date_sort_key(value):
    text = str(value or "").strip()
    candidates = (
        "%Y-%m-%d", "%Y/%m/%d", "%m/%d/%Y", "%d/%m/%Y",
        "%m-%d-%Y", "%d-%m-%Y", "%b %d, %Y", "%B %d, %Y",
    )
    for pattern in candidates:
        try:
            return (0, datetime.strptime(text[:len(pattern) + 4], pattern).date().isoformat())
        except (ValueError, TypeError):
            continue
    try:
        return (0, datetime.fromisoformat(text.replace("Z", "+00:00")).date().isoformat())
    except (ValueError, TypeError):
        return (1, text.casefold())


def _is_total_row(row):
    title = _norm(str(row.get("video_title") or ""))
    return title in {"total", "overall", "all videos", "grand total"}


def _breakdown_recommendations(records):
    dimension = next(
        (
            field for field in (
                "traffic_source", "search_term", "geography",
                "age_group", "gender", "device_type",
            )
            if any(row.get(field) for row in records)
        ),
        None,
    )
    if not dimension:
        return []
    metric = next(
        (
            field for field in ("views", "watch_time_minutes", "viewer_percentage")
            if any(isinstance(row.get(field), (int, float)) for row in records)
        ),
        None,
    )
    if not metric:
        return [{
            "priority": "Context",
            "title": "A viewer/source breakdown was found, but no comparable metric was recognized",
            "evidence": f"The report contains {dimension.replace('_', ' ')} values without a recognized views, watch-time or viewer-percentage column.",
            "action": "Re-export the report with Views or Watch time selected, keeping the same breakdown.",
        }]
    groups = defaultdict(float)
    for row in records:
        label = row.get(dimension)
        value = row.get(metric)
        if label and isinstance(value, (int, float)):
            groups[str(label)] += value
    ranked = sorted(groups.items(), key=lambda item: item[1], reverse=True)
    total = sum(value for _, value in ranked)
    if not ranked or total <= 0:
        return []
    label, value = ranked[0]
    share = value / total * 100
    metric_label = {
        "views": "reported views",
        "watch_time_minutes": "reported watch time",
        "viewer_percentage": "reported viewer percentage",
    }[metric]
    return [{
        "priority": "Context",
        "title": f"Highest reported {dimension.replace('_', ' ')}: {label}",
        "evidence": (
            f"{label} has {value:,.1f} {metric_label} ({share:.1f}% of values "
            f"within this exported breakdown)."
        ),
        "action": (
            "Use this as a research lead: inspect which videos, topics or audience needs "
            "are associated with this segment. This report alone does not prove that the "
            "segment caused performance or that it represents every viewer."
        ),
    }]


def _daily_recommendations(records):
    if any(row.get("video_title") and not _is_total_row(row) for row in records):
        records = [row for row in records if not _is_total_row(row)]
    by_day = defaultdict(lambda: {"views": 0.0, "watch_time_minutes": 0.0})
    has_views = False
    has_watch = False
    for row in records:
        raw_date = row.get("date")
        if not raw_date:
            continue
        day = raw_date.strip()
        if not re.search(r"\d{4}", day):
            continue
        for field in ("views", "watch_time_minutes"):
            value = row.get(field)
            if isinstance(value, (int, float)):
                by_day[day][field] += value
                if field == "views":
                    has_views = True
                if field == "watch_time_minutes":
                    has_watch = True
    days = sorted(by_day, key=_date_sort_key)
    recommendations = []
    if len(days) >= 6 and has_views:
        split = max(1, len(days) // 2)
        earlier = days[:split]
        later = days[split:]
        left_total = sum(by_day[day]["views"] for day in earlier)
        right_total = sum(by_day[day]["views"] for day in later)
        left = left_total / len(earlier) if earlier else 0
        right = right_total / len(later) if later else 0
        if left > 0 and later:
            delta = (right - left) / left * 100
            direction = "increased" if delta > 0 else "decreased"
            recommendations.append({
                "priority": "Context",
                "title": f"Average daily views {direction} {abs(delta):.1f}% between the two halves",
                "evidence": f"Average daily views were {left:,.1f} in the earlier half ({left_total:,.0f} total) and {right:,.1f} in the later half ({right_total:,.0f} total) across {len(days)} reported dates.",
                "action": "Compare uploads and traffic sources across those periods before attributing the change to any title, thumbnail, or publishing decision.",
            })
    if has_watch and len(days) >= 2 and not recommendations:
        recommendations.append({
            "priority": "Context",
            "title": "Daily watch-time data is available",
            "evidence": f"The export contains watch time across {len(days)} dates.",
            "action": "Compare watch time alongside views; more views do not necessarily mean more total watch time.",
        })
    return recommendations


def analyze_studio_export(parsed: dict) -> dict:
    """Return a compact report, grounded actions, and metric qualifications from one parsed export."""
    records = parsed["records"]
    videos = _video_summaries(records)
    fields = set(parsed["fields"])
    recommendations = []
    if parsed["report_type"] == "retention":
        recommendations = _retention_recommendations(records)
    elif parsed["report_type"] == "video_performance" or videos:
        recommendations = _performance_recommendations(videos)
    elif parsed["report_type"] in {"traffic_sources", "breakdown"}:
        recommendations = _breakdown_recommendations(records)
    elif "date" in fields:
        recommendations = _daily_recommendations(records)

    numeric_summary = {}
    summary_records = records
    if any(row.get("video_title") and not _is_total_row(row) for row in records):
        summary_records = [row for row in records if not _is_total_row(row)]
    for display_name, key in (
        ("views", "views"),
        ("watch_time_minutes", "watch_time_minutes"),
        ("impressions", "impressions"),
        ("likes", "likes"),
        ("comments", "comments"),
        ("shares", "shares"),
        ("subscribers_gained", "subscribers_gained"),
        ("subscribers_lost", "subscribers_lost"),
    ):
        values = [row.get(key) for row in summary_records if isinstance(row.get(key), (int, float))]
        if values:
            numeric_summary[display_name] = sum(values)

    ctr_all = [r for r in videos if isinstance(r.get("ctr_percent"), (int, float))]
    ctr_qualified = [
        r for r in ctr_all
        if isinstance(r.get("impressions"), (int, float)) and r["impressions"] >= 50
    ]
    retention_all = [r for r in videos if isinstance(r.get("average_view_percentage"), (int, float))]
    views_all = [r for r in videos if isinstance(r.get("views"), (int, float)) and r["views"] > 0]

    diagnostics = {
        "record_count": len(records),
        "video_count": len(videos),
        "fields_detected": parsed["fields"],
        "ignored_headers": parsed["ignored_headers"],
        "qualifications": {
            "ctr": {
                "has_column": "ctr" in fields,
                "total_with_ctr": len(ctr_all),
                "total_with_impressions": len([r for r in videos if isinstance(r.get("impressions"), (int, float))]),
                "qualified_videos": len(ctr_qualified),
                "min_impressions_required": 50,
                "min_videos_required": 2,
                "met": len(ctr_qualified) >= 2,
                "status": "Qualified" if len(ctr_qualified) >= 2 else (
                    f"Only {len(ctr_qualified)} video(s) met the 50-impression threshold (min. 2 required)"
                    if "ctr" in fields
                    else "Missing column (Impressions click-through rate)"
                ),
            },
            "retention": {
                "has_column": "average_view_percentage" in fields,
                "qualified_videos": len(retention_all),
                "min_videos_required": 2,
                "met": len(retention_all) >= 2,
                "status": "Qualified" if len(retention_all) >= 2 else (
                    f"Only {len(retention_all)} video(s) found (min. 2 required)"
                    if "average_view_percentage" in fields
                    else "Missing column (Average percentage viewed)"
                ),
            },
            "views": {
                "has_column": "views" in fields,
                "qualified_videos": len(views_all),
                "min_videos_required": 2,
                "met": len(views_all) >= 2,
                "status": "Qualified" if len(views_all) >= 2 else (
                    f"Only {len(views_all)} video(s) with positive views found (min. 2 required)"
                    if "views" in fields
                    else "Missing column (Views)"
                ),
            },
        },
    }

    return {
        "filename": parsed["filename"],
        "source_sheet": parsed.get("source_sheet"),
        "report_type": parsed["report_type"],
        "header_row": parsed["header_row"],
        "headers_detected": parsed["header_labels"],
        "ignored_headers": parsed["ignored_headers"],
        "fields": parsed["fields"],
        "record_count": len(records),
        "video_count": len(videos),
        "summary": numeric_summary,
        "video_rows": videos,
        "retention_rows": [
            row for row in records
            if isinstance(row.get("elapsed_seconds"), (int, float))
            or isinstance(row.get("elapsed_percent"), (int, float))
        ],
        "recommendations": recommendations,
        "diagnostics": diagnostics,
        "rows": records,
    }


def title_evidence_key(title: str) -> str:
    """Return the stable, punctuation-normalized key used for title-only matches."""
    return "title:" + _norm(title)


def _retention_curves_by_video(report):
    curves = defaultdict(list)
    for row in report.get("rows", []):
        if not isinstance(row.get("retention_percent"), (int, float)) and not isinstance(
            row.get("relative_retention_index"), (int, float)
        ):
            continue
        point = {
            "elapsed": row.get("elapsed_label"),
            "elapsed_seconds": row.get("elapsed_seconds"),
            "elapsed_percent": row.get("elapsed_percent"),
            "retention_percent": row.get("retention_percent"),
            "relative_retention_index": row.get("relative_retention_index"),
        }
        keys = []
        if row.get("video_id"):
            keys.append(f"id:{row['video_id']}")
        if row.get("video_title"):
            keys.append(title_evidence_key(row["video_title"]))
        for key in keys:
            curves[key].append(point)
    return curves


def build_evidence_index(reports: list[dict]) -> dict:
    """Build a session index without merging metrics from potentially overlapping report periods."""
    index = {}
    for report in reports:
        filename = report.get("filename")
        curves = _retention_curves_by_video(report)
        for row in report.get("video_rows", []):
            keys = []
            if row.get("video_id"):
                keys.append(f"id:{row['video_id']}")
            if row.get("video_title"):
                keys.append(title_evidence_key(row["video_title"]))
            matched_curve = next((curves[key] for key in keys if key in curves), [])
            report_recommendations = report.get("recommendations", [])
            if report.get("report_type") == "retention":
                matching_recommendations = report_recommendations[:5]
            else:
                matching_recommendations = [
                    recommendation for recommendation in report_recommendations
                    if (
                        (row.get("video_id") and recommendation.get("video_id") == row.get("video_id"))
                        or (
                            row.get("video_title")
                            and recommendation.get("video_title")
                            and _norm(row["video_title"]) == _norm(recommendation["video_title"])
                        )
                    )
                ][:5]
            evidence = {
                "source_file": filename,
                "report_type": report.get("report_type"),
                "views": row.get("views"),
                "watch_time_minutes": row.get("watch_time_minutes"),
                "average_view_duration_seconds": row.get("average_view_duration_seconds"),
                "average_view_percentage": row.get("average_view_percentage"),
                "impressions": row.get("impressions"),
                "ctr_percent": row.get("ctr_percent"),
                "subscribers_gained": row.get("subscribers_gained"),
                "subscribers_lost": row.get("subscribers_lost"),
                "date_start": row.get("date_start"),
                "date_end": row.get("date_end"),
                "retention_curve": matched_curve[:80],
                "recommendations": matching_recommendations,
            }
            for key in keys:
                index.setdefault(key, [])
                index[key].append(evidence)
    return index


def add_manual_retention_evidence(index: dict, report: dict, reference: str) -> dict:
    """Attach a retention-only file to an explicitly supplied video ID, URL, or exact title."""
    value = (reference or "").strip()
    video_id = None
    match = re.search(r"(?:v=|youtu\.be/|/shorts/|/embed/)([A-Za-z0-9_-]{11})", value)
    if re.fullmatch(r"[A-Za-z0-9_-]{11}", value):
        video_id = value
    elif match:
        video_id = match.group(1)
    elif "://" in value:
        raise StudioExportError("The supplied video URL did not contain a recognizable video ID.")
    title = None if video_id else value.strip('"\'“”')
    if not video_id and not title:
        raise StudioExportError("Enter a valid video URL/ID or the exact video title to attach this report.")

    curves = []
    for row in report.get("retention_rows", []):
        if not isinstance(row.get("retention_percent"), (int, float)) and not isinstance(
            row.get("relative_retention_index"), (int, float)
        ):
            continue
        curves.append({
            "elapsed": row.get("elapsed_label"),
            "elapsed_seconds": row.get("elapsed_seconds"),
            "elapsed_percent": row.get("elapsed_percent"),
            "retention_percent": row.get("retention_percent"),
            "relative_retention_index": row.get("relative_retention_index"),
        })
    if not curves:
        raise StudioExportError("This report has no recognizable retention curve to attach.")

    evidence = {
        "source_file": report.get("filename"),
        "report_type": report.get("report_type"),
        "date_start": None,
        "date_end": None,
        "retention_curve": curves[:80],
        "recommendations": report.get("recommendations", [])[:5],
    }
    key = f"id:{video_id}" if video_id else title_evidence_key(title)
    index.setdefault(key, [])
    index[key].append(evidence)
    return index


def find_video_evidence(index: dict, video_id: str | None, title: str | None) -> list[dict]:
    """Return de-duplicated source records matched by exact ID first, title second."""
    matches = []
    keys = [f"id:{video_id}"] if video_id else []
    if title:
        keys.append(title_evidence_key(title))
    seen = set()
    for key in keys:
        for item in index.get(key, []):
            identity = (
                item.get("source_file"), item.get("date_start"),
                item.get("date_end"), item.get("report_type"),
            )
            if identity not in seen:
                seen.add(identity)
                matches.append(item)
    return matches


def analyze_uploaded_files(file_payloads: list[tuple[str, bytes]]) -> dict:
    """Analyze uploaded files in-memory and return report structures."""
    reports = []
    for filename, data in file_payloads:
        try:
            parsed = parse_studio_export(filename, data)
            analyzed = analyze_studio_export(parsed)
            reports.append(analyzed)
        except Exception:
            pass
    return {"reports": reports}


def build_video_evidence_index(reports: list[dict]) -> dict:
    """Alias for building evidence index from report structures."""
    return build_evidence_index(reports)


def index_uploaded_studio_files(files) -> tuple[dict, list[str]]:
    """Parse multiple uploaded files in memory and build evidence index."""
    reports = []
    notices = []
    for f in files:
        name = getattr(f, "name", "uploaded_file")
        data = f.getvalue() if hasattr(f, "getvalue") else f
        try:
            parsed = parse_studio_export(name, data)
            analyzed = analyze_studio_export(parsed)
            reports.append(analyzed)
            notices.append(f"Parsed {name}")
        except Exception as e:
            notices.append(f"Failed {name}: {str(e)}")
    return build_evidence_index(reports), notices

